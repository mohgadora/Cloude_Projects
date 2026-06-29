"""Document generation: fill templates with variable values and export to Word."""
import io
import re
from typing import Any

from docx import Document as DocxDocument
from docx.shared import Pt
from sqlalchemy.orm import Session

from models import Document, Template


VARIABLE_PATTERN = re.compile(r"\{\{(\w+)\}\}")


def render_content(content: str, values: dict[str, Any]) -> str:
    """Replace {{variable_key}} placeholders with provided values."""
    def replacer(match: re.Match) -> str:
        key = match.group(1)
        return str(values.get(key, f"[{key}]"))

    return VARIABLE_PATTERN.sub(replacer, content)


def build_document_content(template: Template, values: dict[str, Any]) -> str:
    return render_content(template.content, values)


def export_to_word(title: str, content: str) -> bytes:
    """Generate an in-memory .docx file from plain text content."""
    doc = DocxDocument()

    heading = doc.add_heading(title, level=1)
    heading.runs[0].font.size = Pt(16)

    for paragraph in content.split("\n"):
        p = doc.add_paragraph(paragraph)
        p.runs[0].font.size = Pt(12) if p.runs else None

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def get_document_or_404(
    doc_id: int, tenant_id: int, db: Session
) -> Document:
    doc = (
        db.query(Document)
        .filter(
            Document.id == doc_id,
            Document.tenant_id == tenant_id,
            Document.is_deleted == False,
        )
        .first()
    )
    if doc is None:
        from fastapi import HTTPException, status
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="المستند غير موجود")
    return doc
